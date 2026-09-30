#!/usr/bin/env python3
"""Lê um arquivo CSV e gera um arquivo XLSX com os dados e um gráfico."""

from __future__ import annotations

import argparse
import csv
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import Reference, ScatterChart, Series
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


def detectar_configuracao(caminho: Path, codificacao: str) -> csv.Dialect:
    """Detecta o separador e outras características básicas do CSV."""
    with caminho.open("r", encoding=codificacao, newline="") as arquivo:
        amostra = arquivo.read(8192)

    try:
        return csv.Sniffer().sniff(amostra, delimiters=",;\t|")
    except csv.Error:
        return csv.excel


def converter_valor(valor: str) -> str | int | float | None:
    """Converte textos numéricos para números, preservando os demais textos."""
    texto = valor.strip()
    if texto == "":
        return None

    candidato = texto.replace(",", ".")
    try:
        numero = float(candidato)
        return int(numero) if numero.is_integer() else numero
    except ValueError:
        return texto


def criar_grafico(planilha, numero_linhas: int, numero_colunas: int) -> None:
    """Cria um gráfico usando a primeira coluna como X e a segunda como Y."""
    if numero_linhas < 2 or numero_colunas < 2:
        return

    valores_x = [planilha.cell(linha, 1).value for linha in range(2, numero_linhas + 1)]
    valores_y = [planilha.cell(linha, 2).value for linha in range(2, numero_linhas + 1)]

    if not all(isinstance(v, (int, float)) for v in valores_x + valores_y):
        return

    grafico = ScatterChart()
    grafico.title = "Gráfico dos dados do CSV"
    grafico.style = 13
    grafico.scatterStyle = "smooth"
    grafico.y_axis.title = str(planilha.cell(1, 2).value or "Y")
    grafico.x_axis.title = str(planilha.cell(1, 1).value or "X")
    grafico.height = 10
    grafico.width = 19
    grafico.legend = None

    valores_y_ref = Reference(
        planilha,
        min_col=2,
        min_row=2,
        max_row=numero_linhas,
    )
    valores_x_ref = Reference(
        planilha,
        min_col=1,
        min_row=2,
        max_row=numero_linhas,
    )
    serie = Series(
        valores_y_ref,
        valores_x_ref,
        title=str(planilha.cell(1, 2).value or "Y"),
    )
    grafico.series.append(serie)
    grafico.series[0].graphicalProperties.line.solidFill = "1565C0"
    grafico.series[0].graphicalProperties.line.width = 25000
    grafico.series[0].marker.symbol = "none"

    planilha_grafico = planilha.parent.create_sheet("Gráfico")
    planilha_grafico.sheet_view.showGridLines = False
    planilha_grafico.add_chart(grafico, "A1")
    planilha_grafico.page_setup.orientation = "landscape"
    planilha_grafico.page_setup.fitToWidth = 1
    planilha_grafico.page_setup.fitToHeight = 1
    planilha_grafico.sheet_properties.pageSetUpPr.fitToPage = True
    planilha_grafico.print_area = "A1:K22"


def csv_para_xlsx(
    arquivo_csv: Path,
    arquivo_xlsx: Path | None = None,
    codificacao: str = "latin-1",
) -> Path:
    """Converte o CSV informado em XLSX e devolve o caminho da saída."""
    if not arquivo_csv.is_file():
        raise FileNotFoundError(f"Arquivo não encontrado: {arquivo_csv}")
    if arquivo_csv.suffix.lower() != ".csv":
        raise ValueError("O arquivo de entrada precisa ter a extensão .csv.")

    saida = arquivo_xlsx or arquivo_csv.with_suffix(".xlsx")
    if saida.suffix.lower() != ".xlsx":
        saida = saida.with_suffix(".xlsx")
    saida.parent.mkdir(parents=True, exist_ok=True)

    dialecto = detectar_configuracao(arquivo_csv, codificacao)
    with arquivo_csv.open("r", encoding=codificacao, newline="") as arquivo:
        linhas = [
            [converter_valor(valor) for valor in linha]
            for linha in csv.reader(arquivo, dialecto)
        ]

    if not linhas:
        raise ValueError("O arquivo CSV está vazio.")

    numero_colunas = max(len(linha) for linha in linhas)
    for linha in linhas:
        linha.extend([None] * (numero_colunas - len(linha)))

    pasta_trabalho = Workbook()
    planilha = pasta_trabalho.active
    planilha.title = "Dados"
    planilha.sheet_view.showGridLines = False
    planilha.freeze_panes = "A2"

    for linha in linhas:
        planilha.append(linha)

    preenchimento = PatternFill("solid", fgColor="1565C0")
    for celula in planilha[1]:
        celula.fill = preenchimento
        celula.font = Font(color="FFFFFF", bold=True)
        celula.alignment = Alignment(horizontal="center")

    planilha.auto_filter.ref = planilha.dimensions
    for indice_coluna in range(1, numero_colunas + 1):
        largura = max(
            len(str(planilha.cell(linha, indice_coluna).value or ""))
            for linha in range(1, len(linhas) + 1)
        )
        planilha.column_dimensions[get_column_letter(indice_coluna)].width = min(
            largura + 3, 40
        )

    criar_grafico(planilha, len(linhas), numero_colunas)
    pasta_trabalho.save(saida)
    return saida


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Lê um arquivo CSV e gera um arquivo XLSX."
    )
    parser.add_argument("csv", type=Path, help="arquivo CSV de entrada")
    parser.add_argument(
        "xlsx",
        type=Path,
        nargs="?",
        help="arquivo XLSX de saída (opcional)",
    )
    parser.add_argument(
        "--encoding",
        default="latin-1",
        help="codificação do CSV (padrão: latin-1)",
    )
    argumentos = parser.parse_args()

    try:
        resultado = csv_para_xlsx(
            argumentos.csv,
            argumentos.xlsx,
            argumentos.encoding,
        )
    except (FileNotFoundError, ValueError, UnicodeError, OSError) as erro:
        print(f"Erro: {erro}", file=sys.stderr)
        return 1

    print(f"Arquivo XLSX gerado com sucesso: {resultado.resolve()}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
