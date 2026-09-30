#!/usr/bin/env python3
"""Lê um arquivo XLS antigo e gera um arquivo XLSX moderno correspondente."""

from __future__ import annotations

import argparse
import sys
from pathlib import Path

import xlrd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


def xls_para_xlsx(
    arquivo_xls: Path,
    arquivo_xlsx: Path | None = None,
) -> Path:
    """Converte o XLS informado em XLSX e devolve o caminho da saída."""
    if not arquivo_xls.is_file():
        raise FileNotFoundError(f"Arquivo não encontrado: {arquivo_xls}")
    if arquivo_xls.suffix.lower() != ".xls":
        raise ValueError("O arquivo de entrada precisa ter a extensão .xls.")

    # Define o caminho de saída
    saida = arquivo_xlsx or arquivo_xls.with_suffix(".xlsx")
    if saida.suffix.lower() != ".xlsx":
        saida = saida.with_suffix(".xlsx")
    saida.parent.mkdir(parents=True, exist_ok=True)

    # Abre o arquivo XLS antigo
    pasta_trabalho_xls = xlrd.open_workbook(arquivo_xls)

    # Cria a nova pasta de trabalho XLSX
    pasta_trabalho_xlsx = Workbook()

    # Remove a planilha padrão para evitar uma aba vazia extra
    pasta_trabalho_xlsx.remove(pasta_trabalho_xlsx.active)

    # Itera por todas as abas do XLS
    for indice_planilha in range(pasta_trabalho_xls.nsheets):
        planilha_xls = pasta_trabalho_xls.sheet_by_index(indice_planilha)
        planilha_xlsx = pasta_trabalho_xlsx.create_sheet(title=planilha_xls.name)

        # Configurações visuais iniciais
        planilha_xlsx.sheet_view.showGridLines = False
        if planilha_xls.nrows > 1:
            planilha_xlsx.freeze_panes = "A2"

        # Copia as linhas de uma planilha para a outra[cite: 1]
        for linha in range(planilha_xls.nrows):
            valores_linha = planilha_xls.row_values(linha)
            planilha_xlsx.append(valores_linha)

        # Aplica estilo no cabeçalho e ajusta colunas se houver dados[cite: 1]
        if planilha_xls.nrows > 0:
            preenchimento = PatternFill("solid", fgColor="1565C0")
            for celula in planilha_xlsx[1]:
                celula.fill = preenchimento
                celula.font = Font(color="FFFFFF", bold=True)
                celula.alignment = Alignment(horizontal="center")

            numero_colunas = planilha_xls.ncols
            if numero_colunas > 0:
                planilha_xlsx.auto_filter.ref = planilha_xlsx.dimensions
                for indice_coluna in range(1, numero_colunas + 1):
                    # Calcula a largura ideal para a coluna[cite: 1]
                    largura = max(
                        len(str(planilha_xlsx.cell(linha_idx, indice_coluna).value or ""))
                        for linha_idx in range(1, planilha_xls.nrows + 1)
                    )
                    # Define a largura com limites[cite: 1]
                    planilha_xlsx.column_dimensions[get_column_letter(indice_coluna)].width = min(
                        largura + 3, 40
                    )

    # Garante que haja pelo menos uma planilha caso o arquivo original estivesse vazio
    if not pasta_trabalho_xlsx.sheetnames:
        pasta_trabalho_xlsx.create_sheet("Dados")

    # Salva o arquivo final[cite: 1]
    pasta_trabalho_xlsx.save(saida)
    return saida


def main() -> int:
    # Configura o analisador de argumentos[cite: 1]
    parser = argparse.ArgumentParser(
        description="Lê um arquivo XLS antigo e gera um arquivo XLSX moderno."
    )
    parser.add_argument("xls", type=Path, help="arquivo XLS de entrada")
    parser.add_argument(
        "xlsx",
        type=Path,
        nargs="?",
        help="arquivo XLSX de saída (opcional)",
    )
    argumentos = parser.parse_args()

    # Tenta executar a conversão e trata erros[cite: 1]
    try:
        resultado = xls_para_xlsx(argumentos.xls, argumentos.xlsx)
    except (FileNotFoundError, ValueError, OSError) as erro:
        print(f"Erro: {erro}", file=sys.stderr)
        return 1

    # Retorna sucesso[cite: 1]
    print(f"Arquivo XLSX gerado com sucesso: {resultado.resolve()}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())